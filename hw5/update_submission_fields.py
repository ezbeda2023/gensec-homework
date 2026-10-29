"""Add the required submission fields to the editable homework draft."""
from pathlib import Path

from docx import Document

root = Path(__file__).resolve().parent
path = root / 'submission/hw5-ezbeda2023-DRAFT.docx'
doc = Document(path)
anchor = next(p for p in doc.paragraphs if p.text.startswith('DRAFT -'))
anchor.insert_paragraph_before('Submission information', 'Heading 1')
p = anchor.insert_paragraph_before()
p.add_run('Repository URL: ').bold = True
p.add_run('https://github.com/ezbeda2023/gensec-homework')
recording = (root / 'screencast_url.txt').read_text(encoding='utf-8').strip()
p = anchor.insert_paragraph_before()
p.add_run('Screencast URL: ').bold = True
p.add_run(recording or '[Enter your published screencast URL here]')
doc.add_heading('Final submission steps', level=1)
for text in [
    'Commit and push the latest homework code and submission files to the course repository on GitHub.',
    'Verify that the Repository URL field above points to the course repository.',
    'Enter the published recording URL in the Screencast URL field above and in hw5/screencast_url.txt.',
    'Complete the attack evidence and screenshots, then save this Word document and export it to PDF using File > Export > Create PDF/XPS in Word.',
    'Review the exported PDF for readable text and screenshots, then upload it to the Homework Canvas page before the course deadline.',
]:
    doc.add_paragraph(text, style='List Number')
doc.save(path)
check = Document(path)
assert any(p.text.startswith('Repository URL: https://github.com/') for p in check.paragraphs)
assert any(p.text.startswith('Screencast URL:') for p in check.paragraphs)
print(path)
