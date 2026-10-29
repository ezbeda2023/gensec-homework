"""Recreate the current PDF draft as editable Word paragraphs."""
from pathlib import Path

import pdfplumber
from docx import Document
from docx.shared import Inches, Pt, RGBColor

folder = Path(__file__).resolve().parent / 'submission'
document = Document()
section = document.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.65)
section.left_margin = section.right_margin = Inches(0.75)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    style = document.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
document.styles['Normal'].font.size = Pt(11)
document.styles['Normal'].paragraph_format.space_after = Pt(9)
document.styles['Normal'].paragraph_format.line_spacing = 1.12
document.styles['Title'].font.size = Pt(22)
document.styles['Heading 1'].font.size = Pt(17)
document.styles['Heading 2'].font.size = Pt(14)
headings = {'Target and scope': 'Heading 1', 'Legitimate behavior': 'Heading 2',
            'Experiments prepared': 'Heading 1', 'Limitations and remediation': 'Heading 1'}
with pdfplumber.open(folder / 'hw5-ezbeda2023-DRAFT.pdf') as pdf:
    for index, page in enumerate(pdf.pages):
        if index:
            document.add_page_break()
        paragraph = []
        def flush():
            if paragraph:
                document.add_paragraph(' '.join(paragraph))
                paragraph.clear()
        for line in page.extract_text().splitlines():
            if line.startswith('ezbeda2023 | HW5 | DRAFT'):
                continue
            if line.startswith('HW5:'):
                flush()
                document.add_paragraph(line, 'Title')
            elif line in headings:
                flush()
                document.add_paragraph(line, headings[line])
            elif line.startswith(('ezbeda2023 | Security', 'DRAFT -')):
                flush()
                p = document.add_paragraph(line)
                if line.startswith('DRAFT -'):
                    p.runs[0].bold = True
            else:
                paragraph.append(line)
        flush()
document.core_properties.author = 'ezbeda2023'
document.core_properties.title = 'HW5 Adversarial Prompt Engineering'
target = folder / 'hw5-ezbeda2023-DRAFT.docx'
document.save(target)
print(target)
